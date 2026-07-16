import os
import uuid
from datetime import UTC, date, datetime, time, timedelta
from io import BytesIO
from pathlib import Path

import httpx
from fastapi import FastAPI
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.config import Settings
from app.core.security import hash_password
from app.domain.enums import OrderStatus, UserRole
from app.domain.models import ClientProfile, Driver, Order, RouteAssignment, RoutePlan, User
from app.services.dispatch import run_dispatch_job
from app.services.reports import REPORT_KEY, cleanup_reports, generate_report

AppClient = tuple[FastAPI, httpx.AsyncClient]
PASSWORD = "reports-test-password"


async def _create_user(
    app: FastAPI,
    phone: str,
    role: UserRole,
    *,
    full_name: str,
    district: str = "central",
) -> uuid.UUID:
    async with app.state.session_factory() as session:
        user = User(phone=phone, password_hash=hash_password(PASSWORD), role=role)
        session.add(user)
        await session.flush()
        if role == UserRole.client:
            session.add(ClientProfile(user_id=user.id, full_name=full_name))
        elif role == UserRole.driver:
            session.add(
                Driver(
                    user_id=user.id,
                    full_name=full_name,
                    region=district,
                    capacity=4,
                    is_online=True,
                    shift_start=time(7),
                    shift_end=time(20),
                    home_lat=47.10,
                    home_lon=51.90,
                )
            )
        await session.commit()
        return user.id


async def _headers(client: httpx.AsyncClient, phone: str) -> dict[str, str]:
    response = await client.post(
        "/api/v1/auth/login",
        json={"phone": phone, "password": PASSWORD},
    )
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['tokens']['access_token']}"}


async def _create_order(
    app: FastAPI,
    client_id: uuid.UUID,
    service_date: date,
    desired_time: time,
    *,
    twin_group_id: uuid.UUID | None = None,
) -> uuid.UUID:
    async with app.state.session_factory() as session:
        order = Order(
            client_id=client_id,
            created_by=client_id,
            service_date=service_date,
            desired_time=desired_time,
            pickup_addr="Pickup",
            pickup_lat=47.10,
            pickup_lon=51.90,
            dropoff_addr="Dropoff",
            dropoff_lat=47.11,
            dropoff_lon=51.91,
            escort=False,
            seats=1,
            status=OrderStatus.created,
            twin_group_id=twin_group_id,
        )
        session.add(order)
        await session.commit()
        await session.refresh(order)
        return order.id


async def _publish_plan(
    app: FastAPI,
    client: httpx.AsyncClient,
    headers: dict[str, str],
    service_date: date,
    district: str,
) -> RoutePlan:
    created = await client.post(
        "/api/v1/dispatch/jobs",
        json={"service_date": service_date.isoformat(), "district": district},
        headers=headers,
    )
    assert created.status_code == 202, created.text
    job_id = uuid.UUID(created.json()["job_id"])
    assert await run_dispatch_job(
        {"redis": app.state.redis, "session_factory": app.state.session_factory},
        str(job_id),
    )
    async with app.state.session_factory() as session:
        plan: RoutePlan = (
            await session.scalars(select(RoutePlan).where(RoutePlan.routing_job_id == job_id))
        ).one()
    published = await client.post(f"/api/v1/plans/{plan.id}/publish", headers=headers)
    assert published.status_code == 200, published.text
    return plan


def _use_reports_dir(app: FastAPI, reports_dir: Path) -> Settings:
    settings: Settings = app.state.settings
    overridden = settings.model_copy(update={"reports_dir": str(reports_dir)})
    app.state.settings = overridden
    return overridden


async def _generate_dispatcher_sheet(
    app: FastAPI,
    client: httpx.AsyncClient,
    headers: dict[str, str],
    settings: Settings,
    plan_id: uuid.UUID,
) -> tuple[uuid.UUID, httpx.Response]:
    created = await client.post(
        "/api/v1/reports",
        json={"type": "dispatcher_sheet", "params": {"plan_id": str(plan_id)}},
        headers=headers,
    )
    assert created.status_code == 202, created.text
    report_id = uuid.UUID(created.json()["report_id"])
    assert await generate_report(
        {
            "redis": app.state.redis,
            "session_factory": app.state.session_factory,
            "reports_dir": settings.reports_dir,
        },
        str(report_id),
    )
    status = await client.get(f"/api/v1/reports/{report_id}", headers=headers)
    assert status.status_code == 200, status.text
    assert status.json()["status"] == "done"
    assert status.json()["download_url"] == f"/api/v1/reports/{report_id}/download"
    download = await client.get(f"/api/v1/reports/{report_id}/download", headers=headers)
    assert download.status_code == 200, download.text
    return report_id, download


async def test_dispatcher_sheet_full_published_plan_cycle(
    app_with_client: AppClient,
    tmp_path: Path,
) -> None:
    app, client = app_with_client
    settings = _use_reports_dir(app, tmp_path)
    service_date = date.today() + timedelta(days=80)
    district = "p8-report-cycle"
    await _create_user(
        app,
        "7108000001",
        UserRole.dispatcher,
        full_name="Report Dispatcher",
    )
    client_id = await _create_user(
        app,
        "7108000002",
        UserRole.client,
        full_name="Анна Клиент",
    )
    await _create_user(
        app,
        "7108000003",
        UserRole.driver,
        full_name="Иван Водитель",
        district=district,
    )
    await _create_order(app, client_id, service_date, time(9, 15))
    headers = await _headers(client, "7108000001")
    plan = await _publish_plan(app, client, headers, service_date, district)

    _, download = await _generate_dispatcher_sheet(
        app, client, headers, settings, plan.id
    )
    workbook = load_workbook(BytesIO(download.content))
    sheet = workbook["Диспетчерский лист"]
    rows = list(sheet.iter_rows(values_only=True))
    driver_row = next(index for index, row in enumerate(rows) if row[0] == "Иван Водитель")
    assert rows[driver_row][1] == "7108000003"
    assert rows[driver_row + 1][0] == "Анна Клиент"
    assert any(row[0] == "Исключения" for row in rows)
    methodology = workbook["Методология"]
    assert ("job_id", str(plan.routing_job_id)) in {
        (row[0], row[1]) for row in methodology.iter_rows(values_only=True)
    }


async def test_dispatcher_sheet_contains_desired_time_not_planned_time(
    app_with_client: AppClient,
    tmp_path: Path,
) -> None:
    app, client = app_with_client
    settings = _use_reports_dir(app, tmp_path)
    service_date = date.today() + timedelta(days=81)
    district = "p8-report-time"
    await _create_user(app, "7108000011", UserRole.dispatcher, full_name="Dispatcher Time")
    client_id = await _create_user(
        app, "7108000012", UserRole.client, full_name="Клиент Время"
    )
    await _create_user(
        app,
        "7108000013",
        UserRole.driver,
        full_name="Водитель Время",
        district=district,
    )
    await _create_order(app, client_id, service_date, time(9, 17))
    headers = await _headers(client, "7108000011")

    created = await client.post(
        "/api/v1/dispatch/jobs",
        json={"service_date": service_date.isoformat(), "district": district},
        headers=headers,
    )
    assert created.status_code == 202, created.text
    job_id = uuid.UUID(created.json()["job_id"])
    assert await run_dispatch_job(
        {"redis": app.state.redis, "session_factory": app.state.session_factory},
        str(job_id),
    )
    async with app.state.session_factory() as session:
        plan = (
            await session.scalars(select(RoutePlan).where(RoutePlan.routing_job_id == job_id))
        ).one()
        assignments = list(
            await session.scalars(
                select(RouteAssignment).where(RouteAssignment.plan_id == plan.id)
            )
        )
        assert assignments
        for assignment in assignments:
            assignment.planned_pickup_at = datetime.combine(
                service_date, time(12, 34), tzinfo=UTC
            )
        await session.commit()
    published = await client.post(f"/api/v1/plans/{plan.id}/publish", headers=headers)
    assert published.status_code == 200, published.text

    _, download = await _generate_dispatcher_sheet(
        app, client, headers, settings, plan.id
    )
    workbook = load_workbook(BytesIO(download.content))
    values = {
        str(cell)
        for row in workbook["Диспетчерский лист"].iter_rows(values_only=True)
        for cell in row
        if cell is not None
    }
    assert "09:17" in values
    assert "12:34" not in values


async def test_dispatcher_sheet_keeps_twins_adjacent_and_marked(
    app_with_client: AppClient,
    tmp_path: Path,
) -> None:
    app, client = app_with_client
    settings = _use_reports_dir(app, tmp_path)
    service_date = date.today() + timedelta(days=82)
    district = "p8-report-twins"
    await _create_user(app, "7108000021", UserRole.dispatcher, full_name="Dispatcher Twins")
    first_client = await _create_user(
        app, "7108000022", UserRole.client, full_name="Первый Близнец"
    )
    second_client = await _create_user(
        app, "7108000023", UserRole.client, full_name="Второй Близнец"
    )
    await _create_user(
        app,
        "7108000024",
        UserRole.driver,
        full_name="Водитель Близнецов",
        district=district,
    )
    twin_group_id = uuid.uuid4()
    await _create_order(
        app, first_client, service_date, time(10), twin_group_id=twin_group_id
    )
    await _create_order(
        app, second_client, service_date, time(10), twin_group_id=twin_group_id
    )
    headers = await _headers(client, "7108000021")
    plan = await _publish_plan(app, client, headers, service_date, district)

    _, download = await _generate_dispatcher_sheet(
        app, client, headers, settings, plan.id
    )
    rows = list(
        load_workbook(BytesIO(download.content))["Диспетчерский лист"].iter_rows(
            values_only=True
        )
    )
    twin_rows = [
        index
        for index, row in enumerate(rows)
        if row[0] in {"Первый Близнец", "Второй Близнец"}
    ]
    assert len(twin_rows) == 2
    assert twin_rows[1] == twin_rows[0] + 1
    assert all(rows[index][5] == "Близнецы" for index in twin_rows)


async def test_report_visibility_and_roles_return_expected_statuses(
    app_with_client: AppClient,
    tmp_path: Path,
) -> None:
    app, client = app_with_client
    settings = _use_reports_dir(app, tmp_path)
    await _create_user(app, "7108000031", UserRole.dispatcher, full_name="Owner")
    await _create_user(app, "7108000032", UserRole.dispatcher, full_name="Other")
    await _create_user(app, "7108000033", UserRole.client, full_name="Client Role")
    owner_headers = await _headers(client, "7108000031")
    other_headers = await _headers(client, "7108000032")
    client_headers = await _headers(client, "7108000033")
    missing_id = uuid.uuid4()

    for suffix in ("", "/download"):
        missing = await client.get(
            f"/api/v1/reports/{missing_id}{suffix}", headers=owner_headers
        )
        assert missing.status_code == 404

    created = await client.post(
        "/api/v1/reports",
        json={"type": "dispatcher_sheet", "params": {"plan_id": str(uuid.uuid4())}},
        headers=owner_headers,
    )
    assert created.status_code == 202, created.text
    report_id = uuid.UUID(created.json()["report_id"])
    for suffix in ("", "/download"):
        foreign = await client.get(
            f"/api/v1/reports/{report_id}{suffix}", headers=other_headers
        )
        assert foreign.status_code == 404
        forbidden = await client.get(
            f"/api/v1/reports/{report_id}{suffix}", headers=client_headers
        )
        assert forbidden.status_code == 403

    assert not await generate_report(
        {
            "redis": app.state.redis,
            "session_factory": app.state.session_factory,
            "reports_dir": settings.reports_dir,
        },
        str(report_id),
    )
    failed = await client.get(f"/api/v1/reports/{report_id}", headers=owner_headers)
    assert failed.status_code == 200, failed.text
    assert failed.json()["status"] == "failed"
    assert failed.json()["error"]


async def test_cleanup_reports_removes_only_stale_file_and_metadata(
    app_with_client: AppClient,
    tmp_path: Path,
) -> None:
    app, _ = app_with_client
    settings = _use_reports_dir(app, tmp_path)
    stale_id = uuid.uuid4()
    fresh_id = uuid.uuid4()
    stale_path = tmp_path / f"{stale_id}.xlsx"
    fresh_path = tmp_path / f"{fresh_id}.xlsx"
    stale_path.touch()
    fresh_path.touch()
    old_timestamp = (datetime.now(UTC) - timedelta(days=31)).timestamp()
    os.utime(stale_path, (old_timestamp, old_timestamp))
    await app.state.redis.hset(
        REPORT_KEY.format(report_id=stale_id), mapping={"status": "done"}
    )
    await app.state.redis.hset(
        REPORT_KEY.format(report_id=fresh_id), mapping={"status": "done"}
    )

    removed = await cleanup_reports(
        {"redis": app.state.redis, "reports_dir": settings.reports_dir}
    )

    assert removed == 1
    assert not stale_path.exists()
    assert not await app.state.redis.exists(REPORT_KEY.format(report_id=stale_id))
    assert fresh_path.exists()
    assert await app.state.redis.exists(REPORT_KEY.format(report_id=fresh_id))
