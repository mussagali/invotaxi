"""Asynchronous XLSX reports and Redis-backed reporting statistics."""

from __future__ import annotations

import asyncio
import json
import os
import uuid
from collections import defaultdict
from collections.abc import Awaitable, Iterable
from datetime import UTC, date, datetime, time, timedelta
from pathlib import Path
from typing import Any, cast

import structlog
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from redis.asyncio import Redis
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.core.config import get_settings
from app.domain.enums import OrderStatus, PlanStatus
from app.domain.models import (
    ClientProfile,
    Driver,
    DriverLunch,
    Order,
    OrderEvent,
    RouteAssignment,
    RoutePlan,
    RoutingJob,
    User,
)
from routing import __version__ as routing_version

logger = structlog.get_logger(__name__)

REPORT_TTL_SECONDS = 31 * 24 * 60 * 60
REPORT_RETENTION = timedelta(days=30)
REPORT_KEY = "report:{report_id}"
REPORT_ERROR_MAX_LENGTH = 500

_GREEN = PatternFill("solid", fgColor="70AD47")
_YELLOW = PatternFill("solid", fgColor="FFD966")
_GRAY = PatternFill("solid", fgColor="D9EAD3")
_THIN = Side(style="thin", color="808080")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _decoded_hash(values: dict[Any, Any]) -> dict[str, str]:
    return {
        (key.decode() if isinstance(key, bytes) else str(key)): (
            value.decode() if isinstance(value, bytes) else str(value)
        )
        for key, value in values.items()
    }


def _decoded(value: Any) -> str:
    return value.decode() if isinstance(value, bytes) else str(value)


def _report_redis(ctx: dict[str, Any]) -> Redis:
    return cast(Redis, ctx.get("telemetry_redis") or ctx["redis"])


async def save_report_metadata(redis: Redis, report_id: uuid.UUID, values: dict[str, str]) -> None:
    key = REPORT_KEY.format(report_id=report_id)
    async with redis.pipeline(transaction=True) as pipe:
        pipe.hset(key, mapping=values)
        pipe.expire(key, REPORT_TTL_SECONDS)
        await pipe.execute()


async def get_report_metadata(redis: Redis, report_id: uuid.UUID) -> dict[str, str]:
    values = await cast(
        Awaitable[dict[Any, Any]], redis.hgetall(REPORT_KEY.format(report_id=report_id))
    )
    return _decoded_hash(values)


def aggregate_unassigned(
    rows: Iterable[tuple[time, str]],
) -> dict[str, dict[str, int]]:
    """Aggregate exception rows into half-hour slot -> district counters."""
    result: defaultdict[str, defaultdict[str, int]] = defaultdict(lambda: defaultdict(int))
    for desired_time, district in rows:
        minute = 30 if desired_time.minute >= 30 else 0
        start = desired_time.replace(minute=minute, second=0, microsecond=0)
        end_minutes = start.hour * 60 + start.minute + 30
        end = f"{(end_minutes // 60) % 24:02d}:{end_minutes % 60:02d}"
        slot = f"{start:%H:%M}-{end}"
        result[slot][district] += 1
    return {
        slot: dict(sorted(districts.items()))
        for slot, districts in sorted(result.items())
    }


async def unassigned_stats(
    session: AsyncSession, date_from: date, date_to: date
) -> dict[str, dict[str, int]]:
    plans = list(
        await session.scalars(
            select(RoutePlan).where(
                RoutePlan.service_date >= date_from,
                RoutePlan.service_date <= date_to,
                RoutePlan.routing_job_id.is_not(None),
            )
        )
    )
    jobs_by_id = {
        job.id: job
        for job in await session.scalars(
            select(RoutingJob).where(
                RoutingJob.id.in_(
                    [plan.routing_job_id for plan in plans if plan.routing_job_id is not None]
                )
            )
        )
    }
    scoped: dict[uuid.UUID, str] = {}
    for plan in plans:
        if plan.routing_job_id is None:
            continue
        job = jobs_by_id.get(plan.routing_job_id)
        for order_id in job.unassigned_order_ids or [] if job is not None else []:
            scoped[order_id] = plan.district
    if not scoped:
        return {}
    orders = list(await session.scalars(select(Order).where(Order.id.in_(scoped))))
    return aggregate_unassigned(
        (order.desired_time, scoped[order.id])
        for order in orders
        if order.status == OrderStatus.exception
    )


def _days(date_from: date, date_to: date) -> Iterable[date]:
    current = date_from
    while current <= date_to:
        yield current
        current += timedelta(days=1)


async def anomaly_counts(
    redis: Redis, date_from: date, date_to: date
) -> dict[uuid.UUID, int]:
    counts: defaultdict[uuid.UUID, int] = defaultdict(int)
    for day in _days(date_from, date_to):
        async for raw_key in redis.scan_iter(match=f"anomaly:*:{day.isoformat()}"):
            key = _decoded(raw_key)
            try:
                driver_id = uuid.UUID(key.split(":", 2)[1])
            except (IndexError, ValueError):
                continue
            value = await redis.get(raw_key)
            if value is not None:
                counts[driver_id] += int(value)
    return dict(counts)


def _style_row(sheet: Any, row: int, fill: PatternFill, *, bold: bool = False) -> None:
    for cell in sheet[row]:
        cell.fill = fill
        cell.border = _BORDER
        cell.font = Font(bold=bold)


def _finish_sheet(sheet: Any) -> None:
    sheet.freeze_panes = "A2"
    for column in range(1, sheet.max_column + 1):
        letter = get_column_letter(column)
        width = max(
            (len(str(sheet.cell(row, column).value or "")) for row in range(1, sheet.max_row + 1)),
            default=0,
        )
        sheet.column_dimensions[letter].width = min(max(width + 2, 10), 60)


async def _exception_reasons(
    session: AsyncSession, order_ids: set[uuid.UUID]
) -> dict[uuid.UUID, str]:
    if not order_ids:
        return {}
    events = list(
        await session.scalars(
            select(OrderEvent)
            .where(
                OrderEvent.order_id.in_(order_ids),
                OrderEvent.to_status == OrderStatus.exception,
            )
            .order_by(OrderEvent.ts)
        )
    )
    result: dict[uuid.UUID, str] = {}
    for item in events:
        reason = item.meta.get("reason")
        if reason:
            result[item.order_id] = str(reason)
    return result


def _ordered_driver_orders(orders: list[Order]) -> list[Order]:
    twin_times: defaultdict[uuid.UUID, list[time]] = defaultdict(list)
    for order in orders:
        if order.twin_group_id is not None:
            twin_times[order.twin_group_id].append(order.desired_time)

    def key(order: Order) -> tuple[time, str, time, str]:
        group = order.twin_group_id
        first_time = min(twin_times[group]) if group is not None else order.desired_time
        group_key = str(group) if group is not None else str(order.id)
        return first_time, group_key, order.desired_time, str(order.id)

    return sorted(orders, key=key)


async def _published_plan(session: AsyncSession, plan_id: uuid.UUID) -> RoutePlan:
    plan = await session.get(RoutePlan, plan_id)
    if plan is None or plan.status != PlanStatus.published:
        raise ValueError("published plan not found")
    return plan


async def _dispatcher_workbook(session: AsyncSession, plan_id: uuid.UUID) -> Workbook:
    plan = await _published_plan(session, plan_id)
    assignments = list(
        await session.scalars(
            select(RouteAssignment)
            .where(RouteAssignment.plan_id == plan.id)
            .order_by(RouteAssignment.driver_id, RouteAssignment.seq)
        )
    )
    assigned_ids = {order_id for row in assignments for order_id in row.order_ids}
    job = await session.get(RoutingJob, plan.routing_job_id) if plan.routing_job_id else None
    exception_ids = set(job.unassigned_order_ids or []) if job is not None else set()
    all_ids = assigned_ids | exception_ids
    orders = {
        order.id: order
        for order in await session.scalars(select(Order).where(Order.id.in_(all_ids)))
    }
    client_ids = {order.client_id for order in orders.values()}
    clients = {
        item.user_id: item
        for item in await session.scalars(
            select(ClientProfile).where(ClientProfile.user_id.in_(client_ids))
        )
    }
    users = {
        item.id: item
        for item in await session.scalars(
            select(User).where(User.id.in_(client_ids | {row.driver_id for row in assignments}))
        )
    }
    driver_ids = {row.driver_id for row in assignments}
    drivers = {
        item.user_id: item
        for item in await session.scalars(select(Driver).where(Driver.user_id.in_(driver_ids)))
    }
    by_driver: defaultdict[uuid.UUID, list[Order]] = defaultdict(list)
    for assignment in assignments:
        by_driver[assignment.driver_id].extend(
            orders[order_id] for order_id in assignment.order_ids if order_id in orders
        )

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Диспетчерский лист"
    sheet.append(["ФИО", "Телефон", "Заявленное время", "Откуда", "Куда", "Пометка"])
    _style_row(sheet, 1, _GRAY, bold=True)
    for driver_id in sorted(by_driver, key=lambda item: drivers[item].full_name):
        driver = drivers[driver_id]
        sheet.append(
            [
                driver.full_name,
                users[driver_id].phone,
                driver.vehicle_model or "",
                driver.plate or "",
                "Водитель",
                "",
            ]
        )
        _style_row(sheet, sheet.max_row, _GREEN, bold=True)
        unique = {order.id: order for order in by_driver[driver_id]}
        for order in _ordered_driver_orders(list(unique.values())):
            client = clients[order.client_id]
            notes = []
            if order.escort or client.needs_escort:
                notes.append("Сопровождение")
            if order.twin_group_id is not None:
                notes.append("Близнецы")
            sheet.append(
                [
                    client.full_name,
                    users[order.client_id].phone,
                    order.desired_time.strftime("%H:%M"),
                    order.pickup_addr or "",
                    order.dropoff_addr or "",
                    ", ".join(notes),
                ]
            )
            _style_row(sheet, sheet.max_row, _YELLOW)

    sheet.append([])
    sheet.append(["Исключения"])
    _style_row(sheet, sheet.max_row, _GRAY, bold=True)
    sheet.append(["ФИО", "Телефон", "Заявленное время", "Откуда", "Куда", "Причина"])
    _style_row(sheet, sheet.max_row, _GRAY, bold=True)
    reasons = await _exception_reasons(session, exception_ids)
    for order in sorted(
        (orders[item] for item in exception_ids if item in orders),
        key=lambda item: (item.desired_time, str(item.id)),
    ):
        client = clients[order.client_id]
        sheet.append(
            [
                client.full_name,
                users[order.client_id].phone,
                order.desired_time.strftime("%H:%M"),
                order.pickup_addr or "",
                order.dropoff_addr or "",
                reasons.get(order.id, order.cancel_reason or "Не размещён в плане"),
            ]
        )
        _style_row(sheet, sheet.max_row, _YELLOW)
    _finish_sheet(sheet)

    methodology = workbook.create_sheet("Методология")
    methodology.append(["Параметр", "Значение"])
    _style_row(methodology, 1, _GRAY, bold=True)
    methodology.append(
        ["config", json.dumps((job.params if job else {}) or {}, ensure_ascii=False)]
    )
    methodology.append(["routing_version", routing_version])
    methodology.append(["job_id", str(plan.routing_job_id or "")])
    for row in range(2, methodology.max_row + 1):
        _style_row(methodology, row, PatternFill())
    _finish_sheet(methodology)
    return workbook


async def _day_summary_workbook(
    session: AsyncSession, redis: Redis, plan_id: uuid.UUID
) -> Workbook:
    plan = await _published_plan(session, plan_id)
    assignments = list(
        await session.scalars(select(RouteAssignment).where(RouteAssignment.plan_id == plan.id))
    )
    by_driver: defaultdict[uuid.UUID, set[uuid.UUID]] = defaultdict(set)
    for assignment in assignments:
        by_driver[assignment.driver_id].update(assignment.order_ids)
    counts = [len(order_ids) for order_ids in by_driver.values()]
    job = await session.get(RoutingJob, plan.routing_job_id) if plan.routing_job_id else None
    exception_count = len(set(job.unassigned_order_ids or [])) if job is not None else 0
    lunches = list(
        await session.scalars(select(DriverLunch).where(DriverLunch.plan_id == plan.id))
    )
    all_anomalies = await anomaly_counts(redis, plan.service_date, plan.service_date)
    anomalies = {
        driver_id: count
        for driver_id, count in all_anomalies.items()
        if driver_id in by_driver
    }
    drivers = {
        item.user_id: item.full_name
        for item in await session.scalars(
            select(Driver).where(Driver.user_id.in_(set(by_driver) | set(anomalies)))
        )
    }

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сводка дня"
    sheet.append(["Показатель", "Значение"])
    _style_row(sheet, 1, _GRAY, bold=True)
    metrics: list[tuple[str, Any]] = [
        ("Дата", plan.service_date.isoformat()),
        ("Район", plan.district),
        ("Активных водителей", len(by_driver)),
        ("Заказов размещено", len({item for values in by_driver.values() for item in values})),
        ("Исключений", exception_count),
        ("Заказов на водителя, min", min(counts, default=0)),
        ("Заказов на водителя, avg", round(sum(counts) / len(counts), 2) if counts else 0),
        ("Заказов на водителя, max", max(counts, default=0)),
        ("Обедов полных", sum(item.is_full for item in lunches)),
        ("Обедов коротких", sum(not item.is_full for item in lunches)),
    ]
    for name, value in metrics:
        sheet.append([name, value])
        _style_row(sheet, sheet.max_row, PatternFill())
    sheet.append([])
    sheet.append(["GPS-аномалии", "Количество"])
    _style_row(sheet, sheet.max_row, _GRAY, bold=True)
    for driver_id, count in sorted(anomalies.items(), key=lambda item: (-item[1], str(item[0]))):
        sheet.append([drivers.get(driver_id, str(driver_id)), count])
        _style_row(sheet, sheet.max_row, PatternFill())
    _finish_sheet(sheet)
    return workbook


async def _heatmap_workbook(
    session: AsyncSession, date_from: date, date_to: date
) -> Workbook:
    stats = await unassigned_stats(session, date_from, date_to)
    districts = sorted({district for values in stats.values() for district in values})
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Неразмещённые"
    sheet.append(["Слот", *districts])
    _style_row(sheet, 1, _GRAY, bold=True)
    for slot, values in stats.items():
        sheet.append([slot, *(values.get(district, 0) for district in districts)])
        _style_row(sheet, sheet.max_row, _YELLOW)
    _finish_sheet(sheet)
    return workbook


async def generate_report(ctx: dict[str, Any], report_id: str) -> bool:
    """ARQ entrypoint: build one workbook and atomically publish its path."""
    parsed_id = uuid.UUID(report_id)
    redis = _report_redis(ctx)
    key = REPORT_KEY.format(report_id=parsed_id)
    metadata = _decoded_hash(await cast(Awaitable[dict[Any, Any]], redis.hgetall(key)))
    if not metadata or metadata.get("status") != "queued":
        return metadata.get("status") == "done"
    await cast(Awaitable[int], redis.hset(key, mapping={"status": "running"}))
    await redis.expire(key, REPORT_TTL_SECONDS)
    reports_dir = Path(str(ctx.get("reports_dir") or get_settings().reports_dir))
    target = reports_dir / f"{parsed_id}.xlsx"
    temporary = reports_dir / f".{parsed_id}.tmp.xlsx"
    try:
        params = json.loads(metadata["params"])
        session_factory: async_sessionmaker[Any] = ctx["session_factory"]
        async with session_factory() as session:
            if metadata["type"] == "dispatcher_sheet":
                workbook = await _dispatcher_workbook(session, uuid.UUID(params["plan_id"]))
            elif metadata["type"] == "day_summary":
                workbook = await _day_summary_workbook(
                    session, redis, uuid.UUID(params["plan_id"])
                )
            elif metadata["type"] == "unassigned_heatmap":
                workbook = await _heatmap_workbook(
                    session, date.fromisoformat(params["from"]), date.fromisoformat(params["to"])
                )
            else:
                raise ValueError("unsupported report type")
        await asyncio.to_thread(_save_workbook, workbook, reports_dir, temporary, target)
        await cast(
            Awaitable[int],
            redis.hset(key, mapping={"status": "done", "path": str(target)}),
        )
        await redis.expire(key, REPORT_TTL_SECONDS)
        return True
    except Exception as exc:
        await asyncio.to_thread(_remove_files, [temporary, target])
        error = f"{type(exc).__name__}: {exc}"[:REPORT_ERROR_MAX_LENGTH]
        await cast(
            Awaitable[int],
            redis.hset(
                key,
                mapping={"status": "failed", "path": "", "error": error},
            ),
        )
        await redis.expire(key, REPORT_TTL_SECONDS)
        logger.exception("report.failed", report_id=report_id, error=error)
        return False


async def cleanup_reports(ctx: dict[str, Any]) -> int:
    """Remove XLSX files older than 30 days together with their metadata."""
    redis = _report_redis(ctx)
    reports_dir = Path(str(ctx.get("reports_dir") or get_settings().reports_dir))
    cutoff = datetime.now(UTC).timestamp() - REPORT_RETENTION.total_seconds()
    stale = await asyncio.to_thread(_stale_report_files, reports_dir, cutoff)
    removed = 0
    for report_id, path in stale:
        await asyncio.to_thread(path.unlink, missing_ok=True)
        await redis.delete(REPORT_KEY.format(report_id=report_id))
        removed += 1
    logger.info("reports.cleaned", count=removed)
    return removed


def _save_workbook(
    workbook: Workbook, reports_dir: Path, temporary: Path, target: Path
) -> None:
    reports_dir.mkdir(parents=True, exist_ok=True)
    workbook.save(temporary)
    os.replace(temporary, target)


def _remove_files(paths: Iterable[Path]) -> None:
    for path in paths:
        try:
            path.unlink(missing_ok=True)
        except OSError as exc:
            logger.warning("report.cleanup_failed", path=str(path), error=str(exc))


def _stale_report_files(reports_dir: Path, cutoff: float) -> list[tuple[uuid.UUID, Path]]:
    if not reports_dir.exists():
        return []
    result: list[tuple[uuid.UUID, Path]] = []
    for path in reports_dir.glob("*.xlsx"):
        if path.stat().st_mtime >= cutoff:
            continue
        try:
            result.append((uuid.UUID(path.stem), path))
        except ValueError:
            continue
    return result
